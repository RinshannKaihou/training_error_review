#!/usr/bin/env python3
"""Parse catalog/FAULT_MASTER_REFERENCE.zh.md → exports/fault_catalog.{csv,xlsx}.

md is the only source of truth. Heading count must equal parsed rows.
"""
from __future__ import annotations

import csv
import os
import re
import sys
from collections import Counter
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from category_map import CATEGORY_ORDER, resolve_category  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(ROOT, "catalog", "FAULT_MASTER_REFERENCE.zh.md")
CSV_PATH = os.path.join(ROOT, "exports", "fault_catalog.csv")
XLSX_PATH = os.path.join(ROOT, "exports", "fault_catalog.xlsx")

COLUMNS = [
    "category",
    "id",
    "name",
    "stage",
    "coverage",
    "confidence",
    "source",
    "discovery_date",
    "mechanism",
    "behavior",
]

HEADING_RE = re.compile(r"^### ([A-Z][A-Z0-9-]*\.\d+) `([^`]+)`", re.MULTILINE)
ENTRY_RE = re.compile(
    r"^### ([A-Z][A-Z0-9-]*\.\d+) `([^`]+)`[^\n]*\n\n([^\n]+)\n\n(.*?)(?=\n### |\Z)",
    re.MULTILINE | re.DOTALL,
)


def extract_field(body: str, field_name: str) -> str:
    pattern = rf"- \*\*{re.escape(field_name)}\*\*[：:]?\s*(.*?)(?=\n- \*\*|\n### |\Z)"
    m = re.search(pattern, body, re.DOTALL)
    return m.group(1).strip() if m else ""


def parse_meta(line: str) -> dict:
    meta = {"stage": "", "coverage": "", "confidence": ""}
    sm = re.search(r"stage:\s*`([^`]+)`", line)
    if sm:
        meta["stage"] = sm.group(1).strip()
    cov = re.search(r"Cov:\s*`([^`]+)`", line)
    if cov:
        raw = cov.group(1).split("（")[0].split("(")[0].strip()
        meta["coverage"] = "variant" if raw.startswith("var:") else raw
    conf = re.search(r"置信度\s*`([^`]+)`", line)
    if conf:
        meta["confidence"] = conf.group(1).strip()
    return meta


def extract_source(text: str) -> str:
    urls = re.findall(r"https?://[^\s）)，。]+", text)
    if urls:
        return urls[0]
    arxiv = re.findall(r"arXiv:(\d{4}\.\d{4,5})", text, re.IGNORECASE)
    if arxiv:
        return f"arXiv:{arxiv[0]}"
    issues = re.findall(r"#\d+", text)
    if issues:
        return issues[0]
    return text[:80].rstrip("；，。") if text.strip() else ""


def parse_entries(text: str) -> list[dict]:
    expected_ids = HEADING_RE.findall(text)
    entries = []
    for m in ENTRY_RE.finditer(text):
        fault_id, name, meta_line, body = m.group(1), m.group(2), m.group(3), m.group(4)
        meta = parse_meta(meta_line)
        source_detail = extract_field(body, "来源")
        entries.append(
            {
                "category": resolve_category(fault_id),
                "id": fault_id,
                "name": name,
                "stage": meta["stage"],
                "coverage": meta["coverage"],
                "confidence": meta["confidence"],
                "source": extract_source(source_detail),
                "discovery_date": extract_field(body, "发现来源"),
                "mechanism": extract_field(body, "机制"),
                "behavior": extract_field(body, "行为效应"),
            }
        )
    if len(entries) != len(expected_ids):
        parsed = {e["id"] for e in entries}
        all_ids = {i for i, _ in expected_ids}
        raise SystemExit(
            f"FATAL: parsed {len(entries)} entries but found {len(expected_ids)} headings.\n"
            f"Dropped ids: {sorted(all_ids - parsed)}\n"
            "Most likely a metadata line wrapped onto a second line."
        )
    return entries


def export_csv(entries, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        w.writerows(entries)
    print(f"CSV exported: {path} ({len(entries)} rows)")


def export_xlsx(entries, path):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not available, skipping xlsx")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fault Catalog"
    header_font = Font(name="Microsoft YaHei", bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    for col, name in enumerate(COLUMNS, 1):
        c = ws.cell(1, col, name)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = thin

    palette = [
        "E9F5FE", "FFF8E1", "FDE9E9", "E8F5E9", "F3E5F5", "FFF3E0",
        "E0F7FA", "FCE4EC", "F1F8E9", "EDE7F6", "FFFDE7", "EFEBE9",
        "E8EAF6", "E0F2F1", "FBE9E7", "F4F4F4",
    ]
    fills = {
        cat: PatternFill("solid", fgColor=palette[i % len(palette)])
        for i, cat in enumerate(CATEGORY_ORDER)
    }
    data_font = Font(name="Microsoft YaHei", size=9)
    for r, e in enumerate(entries, 2):
        for col, name in enumerate(COLUMNS, 1):
            cell = ws.cell(r, col, e.get(name, ""))
            cell.font = data_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin
            if name == "category" and e.get("category") in fills:
                cell.fill = fills[e["category"]]

    widths = {
        "category": 16, "id": 10, "name": 40, "stage": 14,
        "coverage": 10, "confidence": 12, "source": 36,
        "discovery_date": 22, "mechanism": 60, "behavior": 50,
    }
    for col, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(name, 14)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(entries)+1}"
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "训练故障编目汇总"
    ws2["A2"] = "导出时间"
    ws2["B2"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws2["A3"] = "总条目"
    ws2["B3"] = len(entries)
    row = 5
    ws2.cell(row, 1, "按类别")
    row += 1
    counts = Counter(e["category"] for e in entries)
    for cat in CATEGORY_ORDER:
        if cat in counts:
            ws2.cell(row, 1, cat)
            ws2.cell(row, 2, counts[cat])
            row += 1
    row += 1
    ws2.cell(row, 1, "按 stage")
    row += 1
    for st, n in Counter(e["stage"] for e in entries).most_common():
        ws2.cell(row, 1, st)
        ws2.cell(row, 2, n)
        row += 1
    ws2.column_dimensions["A"].width = 28
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print(f"XLSX exported: {path} ({len(entries)} rows)")


def main():
    print("Reading:", MD_PATH)
    with open(MD_PATH, encoding="utf-8") as f:
        text = f.read()
    entries = parse_entries(text)
    print(f"Parsed {len(entries)} entries")
    export_csv(entries, CSV_PATH)
    export_xlsx(entries, XLSX_PATH)


if __name__ == "__main__":
    main()
