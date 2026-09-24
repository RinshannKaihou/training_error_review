#!/usr/bin/env python3
"""只读校验训练故障主文档、索引与导出的一致性。"""
import csv
import re
from collections import Counter
from pathlib import Path

import openpyxl
from tools.export_faults import COLUMNS, parse_entries

ROOT = Path(__file__).resolve().parent.parent


def check_rows(rows):
    for row in rows:
        if re.search(r"(?:\n|^)(?:\s*---\s*|\s*#+\s+)", row["discovery_date"]):
            raise ValueError(f"discovery_date 吞入标题/分隔符: {row['id']}")


def main():
    text = (ROOT / "catalog/FAULT_MASTER_REFERENCE.zh.md").read_text()
    entries = parse_entries(text)
    check_rows(entries)
    ids = [row["id"] for row in entries]
    names = [row["name"] for row in entries]
    if len(set(ids)) != len(ids) or len(set(names)) != len(names):
        raise ValueError("正文故障 ID 或名称重复")
    totals = re.findall(r"^> 2026-\d\d-\d\d：.*\*\*(\d+) 条\*\*。$", text, re.M)
    if not totals or int(totals[-1]) != len(entries):
        raise ValueError("最新滚动汇总总数与正文不一致")
    index = (ROOT / "catalog/index/known_names.txt").read_text().splitlines()
    if set(index) != {row["id"] + "\t" + row["name"] for row in entries}:
        raise ValueError("known_names.txt 与正文不同步")
    section = text.split("## 目录", 1)[1].split("\n## ", 1)[0]
    claimed = dict(re.findall(r"^\| ([A-Z][A-Z0-9-]*) \| [^|\n]+ \| (\d+) \|$", section, re.M))
    actual = Counter(row["id"].split(".", 1)[0] for row in entries)
    if not claimed or {k: int(v) for k, v in claimed.items()} != dict(actual):
        raise ValueError("目录分组统计与正文不同步")
    with (ROOT / "exports/fault_catalog.csv").open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames != COLUMNS:
            raise ValueError("CSV 列名与导出脚本不一致")
        rows = list(reader)
    if rows != entries:
        raise ValueError("CSV 与主文档条目不一致；先运行 tools/export_faults.py")
    book = openpyxl.load_workbook(ROOT / "exports/fault_catalog.xlsx", read_only=True, data_only=True)
    sheet = book["Fault Catalog"]
    xrows = sheet.iter_rows(values_only=True)
    if list(next(xrows)) != COLUMNS:
        raise ValueError("XLSX 列名与导出脚本不一致")
    for pos, (row, expected) in enumerate(zip(xrows, rows), 1):
        if [value if value is not None else "" for value in row] != [expected[k] for k in COLUMNS]:
            raise ValueError(f"XLSX 第 {pos} 行与 CSV 不一致")
    if sheet.max_row != len(rows) + 1:
        raise ValueError("XLSX/CSV 行数不一致")
    book.close()
    print(f"PASS {len(rows)} 条：正文汇总、目录、索引、CSV、XLSX、字段边界一致")


if __name__ == "__main__":
    main()
